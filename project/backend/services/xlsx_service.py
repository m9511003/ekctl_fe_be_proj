from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from models.export import ExportRequest

HEADERS = [
    "순번",
    "파일명",
    "측정주파수(MHz)",
    "측정전력(µW)",
    "측정값(dBm)",
    "측정값(dBµV/m)",
    "Limit(dBµV/m)",
    "마진(dB)",
    "판정",
]


def build(req: ExportRequest) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "RE_Result"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in req.rows:
        ws.append(
            [
                r.no,
                r.file,
                r.freq_mhz,
                r.power_uw,
                r.dbm,
                r.dbuv_m,
                r.limit,
                r.margin,
                r.verdict,
            ]
        )

    footer_row = ws.max_row + 2
    ws.cell(row=footer_row, column=1, value=f"※ CF = {req.cf_db} dB 가정치 적용")

    for col_idx, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
            12, len(header) + 2
        )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
