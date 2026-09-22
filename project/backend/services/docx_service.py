import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.table import _Row

from models.export import ExportRequest

TEMPLATE_PATH = Path(__file__).resolve().parent.parent.parent.parent / "data" / "report_template.docx"

STANDARD = "FCC Part 15 Subpart B Class B (3 m)"
TEST_NAME = "2.4 GHz 대역 방사성 방출(RE) 시험"
FIXED_REMARK = "※ 측정 검출기 RMS(Avg) · Limit 기준 QP — 교육용 참고 판정"


def _replace_in_paragraph(paragraph, mapping: dict[str, str]) -> None:
    full_text = "".join(run.text for run in paragraph.runs)
    if "{{" not in full_text:
        return
    for key, val in mapping.items():
        full_text = full_text.replace("{{" + key + "}}", val)
    if paragraph.runs:
        paragraph.runs[0].text = full_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = full_text


def _replace_everywhere(doc, mapping: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, mapping)


def _find_row_template(doc):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "{{r.no}}" in cell.text:
                    return table, row
    raise ValueError("결과 행 템플릿({{r.no}})을 찾을 수 없습니다.")


def _build_remarks(rows) -> str:
    fail_lines = [
        f"{r.freq_mhz:.3f} MHz Limit {abs(r.margin):.2f} dB 초과"
        for r in rows
        if r.verdict == "FAIL" and r.margin is not None
    ]
    if fail_lines:
        return "\n".join(fail_lines) + "\n" + FIXED_REMARK
    return FIXED_REMARK


def build(req: ExportRequest) -> bytes:
    doc = Document(TEMPLATE_PATH)

    today: date = datetime.now().date()
    scalar_map = {
        "doc_no": f"HCT-RE-{today.year}-0001",
        "test_date": today.isoformat(),
        "test_name": TEST_NAME,
        "tester": req.tester,
        "standard": STANDARD,
        "sample_name": req.sample_name,
        "cf_db": f"{req.cf_db:g}",
        "total": str(req.summary.total),
        "pass_cnt": str(req.summary.pass_count),
        "fail_cnt": str(req.summary.fail),
        "verdict": req.summary.verdict,
        "remarks": _build_remarks(req.rows),
        "reviewer": req.reviewer,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    _replace_everywhere(doc, scalar_map)

    table, template_row = _find_row_template(doc)
    anchor = template_row._tr
    for r in req.rows:
        row_map = {
            "r.no": str(r.no),
            "r.freq_mhz": f"{r.freq_mhz:.3f}",
            "r.power_uw": f"{r.power_uw:.2f}",
            "r.dbm": f"{r.dbm:.2f}",
            "r.dbuv_m": f"{r.dbuv_m:.2f}",
            "r.limit": f"{r.limit:.1f}" if r.limit is not None else "-",
            "r.margin": f"{r.margin:.2f}" if r.margin is not None else "-",
            "r.verdict": r.verdict,
        }
        new_tr = copy.deepcopy(anchor)
        anchor.addprevious(new_tr)
        new_row = _Row(new_tr, table)
        for cell in new_row.cells:
            for paragraph in cell.paragraphs:
                _replace_in_paragraph(paragraph, row_map)

    anchor.getparent().remove(anchor)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
